import io

from openpyxl import load_workbook

from app.core.enums import Role
from app.tests.conftest import OFFICE
from app.tests.test_leave import apply, decide

API = "/api/v1"
Q = "from=2026-10-26&to=2026-10-26"


def _setup_day(world, team):
    world.office()
    for n in ("alice", "bob", "carol"):
        world.punch(team.h[n])


def test_employee_cannot_use_reports(world, team):
    for path in (f"reports/attendance?{Q}", f"reports/export?report=attendance&{Q}", "reports/leave"):
        assert world.client.get(f"{API}/{path}", headers=team.h["alice"]).status_code == 403
    assert world.client.get(f"{API}/reports/attendance").status_code == 401


def test_attendance_report_respects_role_scope(world, team):
    _setup_day(world, team)
    names = lambda h: {r["employee_name"] for r in world.client.get(f"{API}/reports/attendance?{Q}", headers=h).json()["rows"]}
    assert names(team.h["mgr"]) == {"Alice A", "Bob B"}
    assert {"Alice A", "Bob B", "Carol C"} <= names(team.h["hr"])
    assert names(team.h["owner"]) >= names(team.h["hr"])


def test_manager_cannot_widen_scope_with_filters(world, team):
    _setup_day(world, team)
    r = world.client.get(f"{API}/reports/attendance?{Q}&employee_id={team.carol.id}", headers=team.h["mgr"]).json()
    assert r["rows"] == []
    r = world.client.get(f"{API}/reports/attendance?{Q}&manager_id={team.owner.id}", headers=team.h["mgr"]).json()
    assert r["rows"] == []


def test_filters_by_mode_status_department(world, team):
    _setup_day(world, team)
    world.punch(team.h["mgr"], mode="WFH", remark="Home office day")
    rows = world.client.get(f"{API}/reports/attendance?{Q}&mode=WFH", headers=team.h["owner"]).json()["rows"]
    assert [r["employee_name"] for r in rows] == ["Mona Manager"]
    rows = world.client.get(f"{API}/reports/attendance?{Q}&status=CHECKED_IN", headers=team.h["owner"]).json()["rows"]
    assert len(rows) == 4
    rows = world.client.get(f"{API}/reports/attendance?{Q}&department=Nope", headers=team.h["owner"]).json()["rows"]
    assert rows == []


def test_date_range_validation(world, team):
    assert world.client.get(f"{API}/reports/attendance?from=2026-10-26&to=2026-10-01", headers=team.h["hr"]).status_code == 422
    assert world.client.get(f"{API}/reports/attendance?from=2024-01-01&to=2026-10-01", headers=team.h["hr"]).status_code == 422
    assert world.client.get(f"{API}/reports/attendance?from=garbage", headers=team.h["hr"]).status_code == 422
    assert world.client.get(f"{API}/reports/nonsense", headers=team.h["hr"]).status_code == 404
    assert world.client.get(f"{API}/reports/export?report=nonsense", headers=team.h["hr"]).status_code == 422


def test_csv_export_scope_headers_and_audit(world, team):
    _setup_day(world, team)
    r = world.client.get(f"{API}/reports/export?report=attendance&format=csv&{Q}", headers=team.h["mgr"])
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert 'attachment; filename="attendance_2026-10-26_2026-10-26.csv"' == r.headers["content-disposition"]
    text = r.content.decode("utf-8-sig")
    assert "Alice A" in text and "Bob B" in text and "Carol C" not in text
    assert text.splitlines()[0].startswith("Date,Code,Employee")
    logs = world.client.get(f"{API}/audit-logs?action=REPORT_EXPORTED", headers=team.h["owner"]).json()
    assert logs["total"] == 1 and logs["items"][0]["details"]["report"] == "attendance"


def test_xlsx_export_is_valid_workbook(world, team):
    _setup_day(world, team)
    r = world.client.get(f"{API}/reports/export?report=attendance-summary&format=xlsx&{Q}", headers=team.h["hr"])
    assert r.status_code == 200 and r.content[:2] == b"PK"
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert header[:3] == ["Code", "Employee", "Department"]
    assert ws.max_row >= 4


def test_export_neutralises_formula_injection(world, team):
    world.make(Role.EMPLOYEE, "evil@example.com", '=HYPERLINK("http://evil","x")')
    csv_text = world.client.get(f"{API}/reports/export?report=attendance-summary&format=csv&{Q}",
                                headers=team.h["hr"]).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text
    assert ",=HYPERLINK" not in csv_text and '"=HYPERLINK' not in csv_text
    x = world.client.get(f"{API}/reports/export?report=attendance-summary&format=xlsx&{Q}", headers=team.h["hr"])
    ws = load_workbook(io.BytesIO(x.content)).active
    values = [ws.cell(row=i, column=2).value for i in range(2, ws.max_row + 1)]
    assert not any(isinstance(v, str) and v.startswith("=") for v in values)


def test_leave_and_balance_reports(world, team):
    rid = apply(world, team.h["alice"], "CL", "2026-10-28", "2026-10-29").json()["id"]
    decide(world, team.h["mgr"], rid, "APPROVE")
    apply(world, team.h["carol"], "SL", "2026-10-28", "2026-10-28")
    rng = "from=2026-10-01&to=2026-10-31"
    leave = world.client.get(f"{API}/reports/leave?{rng}", headers=team.h["mgr"]).json()
    assert [(r["employee_name"], r["status"], r["days"]) for r in leave["rows"]] == [("Alice A", "APPROVED", 2)]
    all_leave = world.client.get(f"{API}/reports/leave?{rng}&status=PENDING_HR", headers=team.h["hr"]).json()
    assert [r["employee_name"] for r in all_leave["rows"]] == ["Carol C"]
    bal = world.client.get(f"{API}/reports/leave-balances?{rng}", headers=team.h["mgr"]).json()
    alice_cl = next(r for r in bal["rows"] if r["employee_name"] == "Alice A" and r["leave_type"] == "CL")
    assert (alice_cl["allocated"], alice_cl["used"], alice_cl["balance"]) == (12, 2, 10)
    assert {r["employee_name"] for r in bal["rows"]} == {"Alice A", "Bob B"}
    led = world.client.get(f"{API}/reports/leave-ledger?from=2026-10-26&to=2026-10-26", headers=team.h["hr"]).json()
    assert {"ALLOCATION", "DEBIT"} <= {r["txn_type"] for r in led["rows"]}


def test_exceptions_report(world, team):
    world.office()
    world.punch(team.h["alice"], lat=OFFICE[0] + 0.01)
    r = world.client.get(f"{API}/reports/exceptions?{Q}", headers=team.h["mgr"]).json()
    assert len(r["rows"]) == 1 and "OUTSIDE_RADIUS" in r["rows"][0]["detail"]
    assert world.client.get(f"{API}/reports/exceptions?{Q}", headers=team.h["owner"]).json()["rows"]


def test_dashboards_by_role(world, team):
    _setup_day(world, team)
    emp = world.client.get(f"{API}/dashboard", headers=team.h["alice"]).json()
    assert "team" not in emp and emp["me"]["today"]["status"] == "CHECKED_IN"
    assert any(b["code"] == "CL" for b in emp["me"]["balances"])
    mgr = world.client.get(f"{API}/dashboard", headers=team.h["mgr"]).json()
    assert mgr["team"]["headcount"] == 2 and "recent_audit" not in mgr and "headcount_total" not in mgr["team"]
    hr = world.client.get(f"{API}/dashboard", headers=team.h["hr"]).json()
    assert hr["team"]["headcount_total"] == 6 and "recent_audit" not in hr
    owner = world.client.get(f"{API}/dashboard", headers=team.h["owner"]).json()
    assert "recent_audit" in owner


def test_security_headers_and_error_shape(world, team):
    r = world.client.get(f"{API}/auth/me", headers=team.h["alice"])
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["cache-control"] == "no-store"
    bad = world.client.post(f"{API}/auth/login", json={"email": "a"})
    assert bad.status_code == 422 and isinstance(bad.json()["detail"], str)
    assert world.client.get("/api/health").json() == {"status": "ok"}
