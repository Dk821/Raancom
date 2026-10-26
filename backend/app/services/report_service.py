"""Role-scoped reports and CSV/XLSX export. All reports share one builder so the
on-screen table and the downloaded file can never disagree about scope."""
import csv
import io
from datetime import date, datetime, timedelta

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.enums import ReviewStatus, Role
from app.core.permissions import Principal, team_clause
from app.models.attendance import AttendanceEvent
from app.models.employee import Employee
from app.models.leave_ledger import LeaveLedger
from app.models.leave_request import LeaveRequest
from app.models.leave_type import LeaveType
from app.models.user import User
from app.services import attendance_service, leave_service
from app.services.settings_service import get_settings, get_tz, local_today

REPORT_NAMES = ("attendance", "attendance-summary", "leave", "leave-balances", "leave-ledger", "exceptions")
MAX_ROWS = 50_000


def _col(key, label):
    return {"key": key, "label": label}


def parse_range(frm: date | None, to: date | None, settings) -> tuple[date, date]:
    today = local_today(settings)
    to = to or today
    frm = frm or to.replace(day=1)
    if frm > to:
        raise HTTPException(422, "'from' date must not be after 'to' date")
    if (to - frm).days > 366:
        raise HTTPException(422, "Date range cannot exceed 366 days")
    return frm, to


def scoped_employees(db: Session, p: Principal, f: dict) -> list[Employee]:
    if p.role == Role.EMPLOYEE.value:
        raise HTTPException(403, "You do not have access to reports")
    stmt = select(Employee).join(User, User.id == Employee.user_id).where(team_clause(p))
    if f.get("employee_id"):
        stmt = stmt.where(Employee.id == f["employee_id"])
    if f.get("department"):
        stmt = stmt.where(Employee.department == f["department"])
    if f.get("manager_id"):
        stmt = stmt.where(Employee.manager_id == f["manager_id"])
    return list(db.scalars(stmt.order_by(Employee.full_name)))


def _hm(dt: datetime | None, tz) -> str:
    return dt.astimezone(tz).strftime("%H:%M") if dt else ""


def build_report(db: Session, p: Principal, name: str, f: dict) -> dict:
    if name not in REPORT_NAMES:
        raise HTTPException(404, "Unknown report")
    settings = get_settings(db)
    tz = get_tz(settings)
    frm, to = parse_range(f.get("from_date"), f.get("to_date"), settings)
    emps = scoped_employees(db, p, f)
    emp_by_id = {e.id: e for e in emps}
    ids = list(emp_by_id)
    cols: list[dict] = []
    rows: list[dict] = []

    if name == "attendance":
        cols = [_col("date", "Date"), _col("employee_code", "Code"), _col("employee_name", "Employee"),
                _col("department", "Department"), _col("status", "Status"), _col("check_in", "Check-in"),
                _col("check_out", "Check-out"), _col("hours", "Hours"), _col("mode", "Mode"),
                _col("late", "Late"), _col("late_minutes", "Late (min)"), _col("geofence_status", "Geofence"),
                _col("distance_m", "Distance (m)"), _col("review_status", "Review"), _col("leave_type", "Leave"),
                _col("corrected", "Corrected")]
        for r in attendance_service.build_rows(db, emps, frm, to):
            if f.get("mode") and r["mode"] != f["mode"]:
                continue
            if f.get("status") and r["status"] != f["status"]:
                continue
            rows.append({
                "date": r["date"].isoformat(), "employee_code": r["employee_code"], "employee_name": r["employee_name"],
                "department": r["department"] or "", "status": r["status"], "check_in": _hm(r["check_in"], tz),
                "check_out": _hm(r["check_out"], tz),
                "hours": round(r["worked_minutes"] / 60, 2) if r["worked_minutes"] is not None else "",
                "mode": r["mode"] or "", "late": "Yes" if r["late"] else "", "late_minutes": r["late_minutes"] or "",
                "geofence_status": r["geofence_status"] or "", "distance_m": r["distance_m"] if r["distance_m"] is not None else "",
                "review_status": r["review_status"] or "", "leave_type": r["leave_type"] or "",
                "corrected": "Yes" if r["corrected"] else "",
            })
        rows.sort(key=lambda x: (x["date"], x["employee_name"]))

    elif name == "attendance-summary":
        cols = [_col("employee_code", "Code"), _col("employee_name", "Employee"), _col("department", "Department"),
                _col("present", "Present"), _col("half_day", "Half day"), _col("absent", "Absent"),
                _col("leave", "Leave"), _col("holiday", "Holiday"), _col("weekly_off", "Weekly off"),
                _col("incomplete", "Incomplete"), _col("late_count", "Late"), _col("hours", "Hours worked"),
                _col("payable_days", "Payable days")]
        if settings.late_rule_enabled:
            cols.append(_col("late_half_day_deductions", "Late deductions (days)"))
        all_rows = attendance_service.build_rows(db, emps, frm, to)
        by_emp: dict[int, list] = {}
        for r in all_rows:
            by_emp.setdefault(r["employee_id"], []).append(r)
        for e in emps:
            s = attendance_service.summarize(by_emp.get(e.id, []), settings)
            row = {"employee_code": e.employee_code, "employee_name": e.full_name, "department": e.department or "",
                   "hours": round(s["worked_minutes"] / 60, 2), **{k: s[k] for k in (
                       "present", "half_day", "absent", "leave", "holiday", "weekly_off", "incomplete", "late_count", "payable_days")}}
            if settings.late_rule_enabled:
                row["late_half_day_deductions"] = s.get("late_half_day_deductions", 0)
            rows.append(row)

    elif name == "leave":
        cols = [_col("id", "Request #"), _col("employee_code", "Code"), _col("employee_name", "Employee"),
                _col("department", "Department"), _col("leave_type", "Leave type"), _col("start_date", "From"),
                _col("end_date", "To"), _col("days", "Days"), _col("half_day", "Half day"), _col("status", "Status"),
                _col("reason", "Reason"), _col("applied_on", "Applied on"), _col("decided_on", "Decided on")]
        stmt = select(LeaveRequest).where(LeaveRequest.employee_id.in_(ids or [-1]),
                                          LeaveRequest.start_date <= to, LeaveRequest.end_date >= frm)
        if f.get("leave_type_id"):
            stmt = stmt.where(LeaveRequest.leave_type_id == f["leave_type_id"])
        if f.get("status"):
            stmt = stmt.where(LeaveRequest.status == f["status"])
        for r in db.scalars(stmt.order_by(LeaveRequest.start_date.desc(), LeaveRequest.id.desc())):
            e = emp_by_id[r.employee_id]
            rows.append({
                "id": r.id, "employee_code": e.employee_code, "employee_name": e.full_name,
                "department": e.department or "", "leave_type": r.leave_type.code, "start_date": r.start_date.isoformat(),
                "end_date": r.end_date.isoformat(), "days": r.days,
                "half_day": r.half_day_session if r.is_half_day else "", "status": r.status, "reason": r.reason,
                "applied_on": r.created_at.astimezone(tz).strftime("%Y-%m-%d"),
                "decided_on": r.decided_at.astimezone(tz).strftime("%Y-%m-%d") if r.decided_at else "",
            })

    elif name == "leave-balances":
        year = f.get("year") or to.year
        cols = [_col("employee_code", "Code"), _col("employee_name", "Employee"), _col("department", "Department"),
                _col("leave_type", "Leave type"), _col("allocated", "Allocated"), _col("adjusted", "Adjusted"),
                _col("used", "Used"), _col("pending", "Pending"), _col("balance", "Balance"), _col("available", "Available")]
        for e in emps:
            for b in leave_service.balance_summary(db, settings, e.id, year):
                if not b["tracked"]:
                    continue
                if f.get("leave_type_id") and b["leave_type_id"] != f["leave_type_id"]:
                    continue
                rows.append({"employee_code": e.employee_code, "employee_name": e.full_name,
                             "department": e.department or "", "leave_type": b["code"], "allocated": b["allocated"],
                             "adjusted": b["adjusted"], "used": b["used"], "pending": b["pending"],
                             "balance": b["balance"], "available": b["available"]})

    elif name == "leave-ledger":
        cols = [_col("date", "Date"), _col("employee_code", "Code"), _col("employee_name", "Employee"),
                _col("leave_type", "Leave type"), _col("year", "Year"), _col("txn_type", "Transaction"),
                _col("days", "Days"), _col("reference", "Reference"), _col("note", "Note"), _col("actor", "By")]
        types = {t.id: t.code for t in db.scalars(select(LeaveType))}
        lo = datetime.combine(frm, datetime.min.time(), tzinfo=tz)
        hi = datetime.combine(to + timedelta(days=1), datetime.min.time(), tzinfo=tz)
        stmt = select(LeaveLedger).where(LeaveLedger.employee_id.in_(ids or [-1]),
                                         LeaveLedger.created_at >= lo, LeaveLedger.created_at < hi)
        if f.get("leave_type_id"):
            stmt = stmt.where(LeaveLedger.leave_type_id == f["leave_type_id"])
        actors = {u.id: u.employee.full_name for u in db.scalars(select(User)) if u.employee}
        for l in db.scalars(stmt.order_by(LeaveLedger.id.desc()).limit(MAX_ROWS)):
            e = emp_by_id[l.employee_id]
            rows.append({
                "date": l.created_at.astimezone(tz).strftime("%Y-%m-%d %H:%M"), "employee_code": e.employee_code,
                "employee_name": e.full_name, "leave_type": types.get(l.leave_type_id, ""), "year": l.year,
                "txn_type": l.txn_type, "days": l.days,
                "reference": f"{l.reference_type} #{l.reference_id}" if l.reference_type else "",
                "note": l.note or "", "actor": actors.get(l.actor_user_id, ""),
            })

    elif name == "exceptions":
        cols = [_col("date", "Date"), _col("employee_code", "Code"), _col("employee_name", "Employee"),
                _col("kind", "Exception"), _col("detail", "Detail"), _col("mode", "Mode"),
                _col("distance_m", "Distance (m)"), _col("review_status", "Review"), _col("review_note", "Review note")]
        stmt = select(AttendanceEvent).where(
            AttendanceEvent.employee_id.in_(ids or [-1]), AttendanceEvent.attendance_date >= frm,
            AttendanceEvent.attendance_date <= to, AttendanceEvent.review_status.in_(
                [ReviewStatus.FLAGGED.value, ReviewStatus.REVIEWED.value]),
            AttendanceEvent.flag_reasons.isnot(None),
        )
        for ev in db.scalars(stmt):
            e = emp_by_id[ev.employee_id]
            rows.append({
                "date": ev.attendance_date.isoformat(), "employee_code": e.employee_code, "employee_name": e.full_name,
                "kind": "Location/Accuracy", "detail": (ev.flag_reasons or "").replace(",", ", ") + f" ({ev.event_type})",
                "mode": ev.mode, "distance_m": ev.distance_m if ev.distance_m is not None else "",
                "review_status": ev.review_status, "review_note": ev.review_note or "",
            })
        for r in attendance_service.build_rows(db, emps, frm, to):
            if r["status"] == "INCOMPLETE":
                rows.append({"date": r["date"].isoformat(), "employee_code": r["employee_code"],
                             "employee_name": r["employee_name"], "kind": "Incomplete day",
                             "detail": "Checked in without checking out", "mode": r["mode"] or "", "distance_m": "",
                             "review_status": "", "review_note": ""})
        rows.sort(key=lambda x: (x["date"], x["employee_name"]), reverse=True)

    if len(rows) > MAX_ROWS:
        raise HTTPException(422, "Too many rows; narrow the filters")
    return {"name": name, "from": frm.isoformat(), "to": to.isoformat(), "columns": cols, "rows": rows}


# -------------------------------------------------------------------- export
def _safe(v):
    """Neutralise spreadsheet formula injection in text cells."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def export_csv(report: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in report["columns"]])
    for r in report["rows"]:
        w.writerow([_safe(r.get(c["key"], "")) for c in report["columns"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def export_xlsx(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = report["name"][:31]
    ws.append([c["label"] for c in report["columns"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in report["rows"]:
        ws.append([_safe(r.get(c["key"], "")) for c in report["columns"]])
    for i, c in enumerate(report["columns"], start=1):
        width = max([len(str(c["label"]))] + [len(str(r.get(c["key"], ""))) for r in report["rows"][:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 50)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
